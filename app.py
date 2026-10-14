"""Витрина отчета: план / факт / остаток / средняя цена по месяц - магазин - бит литера - категория.

Читает план (csv), Sales_sku_target и Ostatki_target с Google Drive, перезаписывает report.xlsx на Google Drive.
Месяцы: все месяцы плана с января 2026 по месяц последней даты продаж.
- факт: сумма_со_скидкой за месяц (для текущего месяца - MTD по as_of);
- остаток: срез на 1-е число месяца; если среза на 1-е нет (текущий месяц) - последний срез внутри месяца.

Запуск:
    python app.py                  # MTD на последнюю дату в продажах
    python app.py --as-of 2026-10-03
    python app.py --schedule       # висит в фоне и обновляет report.xlsx сразу и далее каждые 24 часа
"""
import argparse
import io
import logging
import os
import textwrap
from datetime import datetime
from pathlib import Path

import numpy as np
import pandas as pd
from apscheduler.schedulers.blocking import BlockingScheduler
from dotenv import load_dotenv
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from drive_client import download_bytes, get_drive_service, update_file_bytes

BASE_DIR = Path(__file__).parent
logger = logging.getLogger("app")

load_dotenv(BASE_DIR / ".env")

CREDS_PATH = BASE_DIR / os.environ["DRIVE_CREDS_PATH"]
TOKEN_PATH = BASE_DIR / os.environ["DRIVE_TOKEN_PATH"]
SALES_SKU_ID = os.environ["SALES_SKU_TARGET_ID"]
OSTATKI_ID = os.environ["OSTATKI_TARGET_ID"]

PLAN_ID = os.environ["PLAN_PRODAZH_ID"]
REPORT_XLSX_ID = os.environ["REPORT_XLSX_ID"]
XLSX_MIMETYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

# форматы Excel: колонка -> (формат, вид отображения для расчета ширины)
CURRENCY = r"[$₸-43F]"  # тенге (kk-KZ)
FINANCIAL = rf'_-* #,##0\ {CURRENCY}_-;\-* #,##0\ {CURRENCY}_-;_-* "-"\ {CURRENCY}_-;_-@_-'  # финансовый с валютой, без десятичных знаков
COLUMN_FORMATS = {
    "доля бит литеры": ("0%", ("percent", 0)),
    "доля категории внутри бит литеры": ("0%", ("percent", 0)),
    "выполнение плана": ("0.00%", ("percent", 2)),
    "план месяц": (FINANCIAL, ("financial", 0)),
    "план": (FINANCIAL, ("financial", 0)),
    "факт_mtd": (FINANCIAL, ("financial", 0)),
    "плана осталось": (FINANCIAL, ("financial", 0)),
    "средняя цена продажи с учетом скидки": (FINANCIAL, ("financial", 0)),
}
HEADER_LINES = 3  # высота строки заголовков = столько обычных строк
ROW_HEIGHT = 15  # высота обычной строки в Excel, pt

# Средняя ежемесячная скидка от розничной цены (память проекта, 2026-10-04)
DISCOUNT = {"ЮИ": 0.10, "ЗИ": 0.10, "СГ": 0.0}

# колонки плана в csv (порядок и названия не меняем)
PLAN_COLUMNS = ["месяц", "магазин", "бит литера", "категория", "доля бит литеры",
                "доля категории внутри бит литеры", "план месяц", "план"]

# остаток / шт продать осталось: от этого значения запас считается достаточным
STOCK_COVER_OK = 3

KEY = ["магазин", "бит_литера", "категория"]


def calc_fact(sales: pd.DataFrame, month_start: pd.Timestamp, as_of: pd.Timestamp) -> pd.Series:
    month_end = month_start + pd.offsets.MonthEnd(0)
    period = sales[(sales["date"] >= month_start) & (sales["date"] <= min(month_end, as_of))]
    return period.groupby(KEY)["сумма_со_скидкой"].sum().rename("факт_mtd")


def pick_snapshot(snapshots: list, month_start: pd.Timestamp) -> pd.Timestamp:
    """Срез на 1-е число месяца; если его нет - последний срез внутри месяца."""
    month_end = month_start + pd.offsets.MonthEnd(0)
    in_month = [d for d in snapshots if month_start <= d <= month_end]
    if not in_month:
        raise SystemExit(f"Нет среза остатков за {month_start:%m-%Y}")
    return month_start if month_start in in_month else max(in_month)


def calc_stock(snapshot: pd.Timestamp, ostatki: pd.DataFrame) -> pd.DataFrame:
    """Остаток шт и средневзвешенная (по остатку) розничная цена на последнем срезе."""
    last = ostatki[ostatki["date"] == snapshot].copy()
    last["сумма_розница"] = last["розничная_цена"] * last["остаток"]
    g = last.groupby(KEY)[["остаток", "сумма_розница"]].sum()
    g["розничная_цена_ср"] = g["сумма_розница"] / g["остаток"].where(g["остаток"] > 0)
    return g.rename(columns={"остаток": "остаток_шт"})[["остаток_шт", "розничная_цена_ср"]]


def read_drive_csv(service, file_id: str, **kwargs) -> pd.DataFrame:
    """Скачивает csv с Google Drive по id и читает его в DataFrame."""
    return pd.read_csv(download_bytes(service, file_id), **kwargs)


def display_length(value, display=None) -> int:
    """Длина значения так, как оно отображается в ячейке."""
    if value is None or value == "" or (isinstance(value, float) and np.isnan(value)):
        return 0
    kind, digits = display or (None, 0)
    if kind == "percent":
        return len(f"{value * 100:.{digits}f}%")
    if kind == "financial":
        return len(f"{value:,.{digits}f}") + 4  # " ₸" и отступы формата _- и -_
    if isinstance(value, float):
        return len(f"{value:.10g}")
    return len(str(value))


def header_width(header: str, min_width: int) -> int:
    """Минимальная ширина, при которой заголовок переносится не более чем на HEADER_LINES строк."""
    width = max(min_width, max(len(word) for word in header.split()))
    while len(textwrap.wrap(header, width)) > HEADER_LINES:
        width += 1
    return width


def format_sheet(ws, columns: list) -> None:
    """Таблица Excel: стиль таблицы, заголовки, форматы чисел, ширина столбцов по значениям."""
    last_col, last_row = get_column_letter(ws.max_column), ws.max_row
    table = Table(displayName="report", ref=f"A1:{last_col}{last_row}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(table)
    ws.row_dimensions[1].height = ROW_HEIGHT * HEADER_LINES

    for idx, name in enumerate(columns, start=1):
        letter = get_column_letter(idx)
        ws[f"{letter}1"].alignment = Alignment(wrap_text=True, vertical="center")
        number_format, display = COLUMN_FORMATS.get(name, (None, None))
        width = 0
        for (cell,) in ws.iter_rows(min_row=2, min_col=idx, max_col=idx):
            if number_format:
                cell.number_format = number_format
            width = max(width, display_length(cell.value, display))
        # запас под фильтр в заголовке; заголовок должен влезать в HEADER_LINES строк
        ws.column_dimensions[letter].width = header_width(name, max(width, 6) + 3)


def build_report_xlsx(df: pd.DataFrame, sheet_name: str) -> bytes:
    """xlsx в памяти: один лист с отформатированной таблицей (общий для витрины и бота)."""
    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name=sheet_name, index=False)
        format_sheet(writer.sheets[sheet_name], list(df.columns))
    return buffer.getvalue()


def update_report(as_of_arg: str = None) -> None:
    service = get_drive_service(CREDS_PATH, TOKEN_PATH)
    plan = read_drive_csv(service, PLAN_ID, dtype={"месяц": str}).rename(columns={"бит литера": "бит_литера"})

    sales = read_drive_csv(service, SALES_SKU_ID, usecols=["date", *KEY, "сумма_со_скидкой"], parse_dates=["date"])
    as_of = pd.Timestamp(as_of_arg) if as_of_arg else sales["date"].max()

    ostatki = read_drive_csv(service, OSTATKI_ID, usecols=["date", *KEY, "розничная_цена", "остаток"], parse_dates=["date"])
    snapshots = sorted(ostatki["date"].unique())

    parts = []
    # месяцы в отчете от последнего к первому; "месяц" - строка ММ-ГГГГ, сортируем по дате, а не по строке
    months = sorted(
        ((pd.Timestamp(f"{month[3:]}-{month[:2]}-01"), plan_m) for month, plan_m in plan.groupby("месяц")),
        key=lambda item: item[0],
        reverse=True,
    )
    for month_start, plan_m in months:
        if month_start > as_of:
            continue
        snapshot = pick_snapshot(snapshots, month_start)
        part = (
            plan_m.merge(calc_fact(sales, month_start, as_of), on=KEY, how="left")
            .merge(calc_stock(snapshot, ostatki), on=KEY, how="left")
        )
        part["дата_факта"] = min(month_start + pd.offsets.MonthEnd(0), as_of).date()
        part["дата_остатка"] = snapshot.date()
        parts.append(part)

    # витрина строится по плану: филиалы и категории без плана в отчет не входят
    out = pd.concat(parts, ignore_index=True)
    out["факт_mtd"] = out["факт_mtd"].fillna(0)
    out["остаток_шт"] = out["остаток_шт"].fillna(0)
    out["выполнение плана"] = out["факт_mtd"] / out["план"].where(out["план"] > 0)
    left = out["план"] - out["факт_mtd"]
    out["плана осталось"] = left.where(left >= 0)  # при перевыполнении плана пусто
    # без остатка (розничная_цена_ср пустая) средняя цена пустая
    price = out["розничная_цена_ср"] * (1 - out["бит_литера"].map(DISCOUNT))
    out["средняя цена продажи с учетом скидки"] = (price / 10).round() * 10  # округление до десятков
    price_col = out["средняя цена продажи с учетом скидки"]
    out["шт продать осталось"] = np.ceil(out["плана осталось"] / price_col.where(price_col > 0))  # вверх до целых
    need, stock = out["шт продать осталось"], out["остаток_шт"]
    enough = stock >= need * STOCK_COVER_OK  # остаток / шт продать >= 3 (при need=0 тоже верно)
    no_stock = (stock == 0) & (out["плана осталось"] > 0)  # цены нет (остаток 0), но план не выполнен
    out["возможность выполнить план"] = np.select(
        [no_stock, need.isna(), stock < need, ~enough],
        ["нет, остатков не хватает", None, "нет, остатков не хватает", "да, рекомендуется пополнить запасы"],
        default="да, достаточный запас",
    )
    out.loc[need.isna() & ~no_stock, "возможность выполнить план"] = np.nan

    # колонки плана остаются как в csv, справа добавляются расчетные
    out = out.rename(columns={"бит_литера": "бит литера"})
    out = out[[*PLAN_COLUMNS, "факт_mtd", "выполнение плана", "плана осталось", "остаток_шт", "средняя цена продажи с учетом скидки", "шт продать осталось", "возможность выполнить план", "дата_факта", "дата_остатка"]]
    sheet_name = f"{out['дата_факта'].max():%d.%m.%Y}"  # бот (/updated) берет дату обновления из названия листа
    update_file_bytes(service, REPORT_XLSX_ID, build_report_xlsx(out, sheet_name), XLSX_MIMETYPE)

    in_period = sales[(sales["date"] >= pd.Timestamp("2026-01-01")) & (sales["date"] <= as_of)]
    print(f"report.xlsx (Drive): {len(out)} строк, месяцев {out['месяц'].nunique()}, факт по {as_of.date()}")
    print(f"Продажи вне плана (не вошли в витрину): {in_period['сумма_со_скидкой'].sum() - out['факт_mtd'].sum():,.0f}")


def run_scheduler() -> None:
    """Обновляет report.xlsx каждые 24 часа. Ошибка в одном запуске не останавливает расписание."""
    scheduler = BlockingScheduler()
    scheduler.add_job(update_report, "interval", hours=24, next_run_time=datetime.now())  # первый запуск сразу
    logger.info("Планировщик запущен: обновление report.xlsx сразу и далее каждые 24 часа")
    scheduler.start()


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--as-of", help="дата, по которую считать MTD (по умолчанию - последняя дата продаж)")
    parser.add_argument("--schedule", action="store_true", help="обновлять report.xlsx каждые 24 часа")
    args = parser.parse_args()
    if args.schedule:
        logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s")
        run_scheduler()
    else:
        update_report(args.as_of)


if __name__ == "__main__":
    main()
