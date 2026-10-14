"""Telegram-бот товарного отчета (long polling).

Команды:
    /updated   - название листа report.xlsx на Google Drive (там записана максимальная дата_факта)
    /report    - отправляет в чат report.xlsx: по файлу на каждый доступный филиал ({филиал}_{as_of}.xlsx),
                 а при доступе all - общий отчет по всем филиалам
    /my_filial - список доступных филиалов

Доступ берется из Google Таблицы Users (колонки user_id, access): user_id - числовой id в Telegram,
access - название филиала (значение в поле "магазин" отчета) или all.

Запуск:
    python bot.py            # слушать команды
    python bot.py --updated  # вывести дату обновления и выйти (без Telegram)
"""
import argparse
import csv
import io
import logging
import os
import time
from collections import defaultdict
from pathlib import Path

import pandas as pd
import requests
from dotenv import load_dotenv
from openpyxl import load_workbook

from app import build_report_xlsx
from drive_client import download_bytes, export_sheet_csv, get_drive_service

BASE_DIR = Path(__file__).parent

load_dotenv(BASE_DIR / ".env")

CREDS_PATH = BASE_DIR / os.environ["DRIVE_CREDS_PATH"]
TOKEN_PATH = BASE_DIR / os.environ["DRIVE_TOKEN_PATH"]
BOT_TOKEN = os.environ["TELEGRAM_BOT_TOKEN"]
REPORT_XLSX_ID = os.environ["REPORT_XLSX_ID"]
USERS_SHEET_ID = os.environ["USERS_SHEET_ID"]
XLSX_MIMETYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

API_URL = f"https://api.telegram.org/bot{BOT_TOKEN}"
POLL_TIMEOUT = 50  # сек, long polling

ACCESS_ALL = "all"
SHOP_COLUMN_IDX = 1  # "магазин" - второй столбец report.xlsx

logger = logging.getLogger("bot")


def get_updated_date() -> str:
    """Дата обновления данных - название листа отчета на Drive (ДД.ММ.ГГГГ)."""
    service = get_drive_service(CREDS_PATH, TOKEN_PATH)
    workbook = load_workbook(download_bytes(service, REPORT_XLSX_ID), read_only=True)
    return workbook.sheetnames[0]


def send_message(chat_id: int, text: str) -> None:
    response = requests.post(f"{API_URL}/sendMessage", json={"chat_id": chat_id, "text": text}, timeout=30)
    response.raise_for_status()


def send_document(chat_id: int, filename: str, data: bytes) -> None:
    response = requests.post(
        f"{API_URL}/sendDocument",
        data={"chat_id": chat_id},
        files={"document": (filename, data, XLSX_MIMETYPE)},
        timeout=120,
    )
    response.raise_for_status()


def get_access(user_id: int) -> list[str]:
    """Филиалы пользователя из таблицы Users (в порядке строк); ['all'] - все филиалы; [] - нет доступа."""
    service = get_drive_service(CREDS_PATH, TOKEN_PATH)
    rows = csv.DictReader(io.StringIO(export_sheet_csv(service, USERS_SHEET_ID)))
    access = defaultdict(list)
    for row in rows:
        access[row["user_id"].strip()].append(row["access"].strip())
    filials = access.get(str(user_id), [])
    if any(f.lower() == ACCESS_ALL for f in filials):
        return [ACCESS_ALL]
    return filials


def filter_report(report: pd.DataFrame, filial: str, as_of: str) -> bytes | None:
    """xlsx (в памяти) только со строками филиала в том же формате, что и report.xlsx; None, если строк нет."""
    filial_rows = report[report.iloc[:, SHOP_COLUMN_IDX] == filial]
    if filial_rows.empty:
        return None
    return build_report_xlsx(filial_rows, as_of)


def handle_updated(chat_id: int, user_id: int) -> None:
    send_message(chat_id, f"Данные обновлены по {get_updated_date()}")


def handle_my_filial(chat_id: int, user_id: int) -> None:
    filials = get_access(user_id)
    if not filials:
        send_message(chat_id, f"Нет доступа. Ваш user_id: {user_id}")
    elif filials == [ACCESS_ALL]:
        send_message(chat_id, "Доступен отчет по всем филиалам")
    else:
        send_message(chat_id, ", ".join(filials))


def handle_report(chat_id: int, user_id: int) -> None:
    filials = get_access(user_id)
    if not filials:
        send_message(chat_id, f"Нет доступа. Ваш user_id: {user_id}")
        return

    service = get_drive_service(CREDS_PATH, TOKEN_PATH)
    report = download_bytes(service, REPORT_XLSX_ID).getvalue()
    as_of = load_workbook(io.BytesIO(report), read_only=True).sheetnames[0]  # дата обновления - название листа

    if filials == [ACCESS_ALL]:
        send_document(chat_id, f"Все филиалы_{as_of}.xlsx", report)
        return

    no_data = []
    report_df = pd.read_excel(io.BytesIO(report))  # читаем один раз, дальше только фильтр по филиалу
    for col in report_df.select_dtypes("datetime"):
        report_df[col] = report_df[col].dt.date  # как в app.py: даты без времени
    for filial in filials:
        filial_report = filter_report(report_df, filial, as_of)
        if filial_report is None:
            no_data.append(filial)
        else:
            send_document(chat_id, f"{filial}_{as_of}.xlsx", filial_report)
    if no_data:
        send_message(chat_id, f"Нет данных в отчете по филиалам: {', '.join(no_data)}")


HANDLERS = {"/updated": handle_updated, "/report": handle_report, "/my_filial": handle_my_filial}


def handle_update(update: dict) -> None:
    message = update.get("message") or {}
    text = message.get("text", "")
    if not text.startswith("/"):
        return
    command = text.split()[0].split("@")[0].lower()  # "/updated@bot_name" -> "/updated"
    handler = HANDLERS.get(command)
    if handler:
        handler(message["chat"]["id"], message["from"]["id"])


def run() -> None:
    offset = None
    logger.info("Бот запущен")
    while True:
        try:
            response = requests.get(
                f"{API_URL}/getUpdates",
                params={"timeout": POLL_TIMEOUT, "offset": offset},
                timeout=POLL_TIMEOUT + 10,
            )
            response.raise_for_status()
            for update in response.json()["result"]:
                offset = update["update_id"] + 1
                try:
                    handle_update(update)
                except Exception:
                    logger.exception("Ошибка обработки update %s", update["update_id"])
        except requests.RequestException:
            logger.exception("Ошибка getUpdates, повтор через 5 сек")
            time.sleep(5)


if __name__ == "__main__":
    logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s")
    logging.getLogger("httpx").setLevel(logging.WARNING)
    parser = argparse.ArgumentParser()
    parser.add_argument("--updated", action="store_true", help="вывести дату обновления и выйти")
    args = parser.parse_args()
    if args.updated:
        print(get_updated_date())
    else:
        run()
